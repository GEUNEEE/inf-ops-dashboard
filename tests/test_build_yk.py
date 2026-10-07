# tests/test_build_yk.py — 영끌러님 대시보드 파서/집계 단위 테스트 (stdlib unittest + openpyxl 픽스처)
# 실행: PYTHONUTF8=1 .venv\Scripts\python.exe -m unittest tests.test_build_yk -v
import sys
import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
for sub in ("excel-parser", "dashboard-builder"):
    sys.path.insert(0, str(BASE / ".claude" / "skills" / sub / "scripts"))

import openpyxl  # noqa: E402

import parse_mail  # noqa: E402
import parse_inf   # noqa: E402
import build_snapshot  # noqa: E402

D = datetime


# ---------------------------------------------------------------- 픽스처 생성
MAIL_HEADERS_OLD = ["분류", "", "", "", "채널명", "채널링크"] + [""] * 19 + \
    ["발송일", "확인일", "회신일", "진행 상태", "미팅일", "협찬 수락일", "광고 수락일"]


def make_mail_wb(path: Path, rows: list, with_owner: bool, sheet="메일발송현황", owner_marks=None):
    """rows: [(send, reply, status, mtg, accept, ad)], owner_marks: 행별 구분 값 리스트(or None)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"] = "미니 대시보드"           # 1~5행은 집계와 무관
    headers = (["구분"] if with_owner else []) + MAIL_HEADERS_OLD
    for c, h in enumerate(headers, start=1):
        ws.cell(row=6, column=c, value=h)
    off = 1 if with_owner else 0
    for i, (send, reply, status, mtg, accept, ad) in enumerate(rows):
        r = 7 + i
        if with_owner and owner_marks is not None:
            ws.cell(row=r, column=1, value=owner_marks[i])
        ws.cell(row=r, column=off + 5, value=f"채널{i}")
        for j, v in enumerate((send, "", reply, status, mtg, accept, ad)):
            ws.cell(row=r, column=off + 26 + j, value=v)
    wb.save(path)


MAIL_ROWS = [
    (D(2026, 9, 1), D(2026, 9, 2), "미팅 대기", D(2026, 9, 5), None, None),
    (D(2026, 9, 3), None, "", None, None, None),
    (D(2026, 9, 4), "필요", "1차 체험진행", None, D(2026, 9, 10), None),
    (D(2026, 8, 20), D(2026, 8, 21), "검토", D(2026, 8, 25), D(2026, 8, 26), None),   # 기타 → 체험 제외
    (D(2026, 8, 22), D(2026, 8, 23), "1차 광고완료", D(2026, 8, 24), D(2026, 8, 25), D(2026, 9, 1)),
    (None, D(2026, 9, 9), "", None, None, None),   # 발송일 없음 → 응답만
]
EXPECTED_ALL = {"total_sent": 5, "etc_excluded": 1, "replied": 5, "meeting_total": 3,
                "exp_total_approx": 2, "ad_total": 1}


def make_master_wb(path: Path, extra_rows_before_name=0, product_row=True, yk_row=True):
    """종합 관리시트 유사 픽스처(전치형). 인플루언서 3명: 가, 나, 다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "인플루언서관리"
    labels = [
        (2, "항목"), (3, "미팅 대기"),
    ]
    r = 9
    if yk_row:
        ws.cell(row=r, column=1, value="영끌러님 관리"); r += 1
    ws.cell(row=r, column=1, value="항목"); r += 1
    for _ in range(extra_rows_before_name):
        ws.cell(row=r, column=1, value="새 항목"); r += 1
    ws.cell(row=r, column=1, value="현재 상태"); r_status = r; r += 1
    ws.cell(row=r, column=1, value="유튜버명"); r_name = r; r += 1
    if product_row:
        ws.cell(row=r, column=1, value="제품"); r_prod = r; r += 1
    ws.cell(row=r, column=1, value="체험 수락일"); r_acc = r; r += 1
    ws.cell(row=r, column=1, value="제품 발송일"); r_ship = r; r += 1
    ws.cell(row=r, column=1, value="1차 광고"); r_ad1 = r; r += 1
    ws.cell(row=r, column=1, value="2차 체험"); r_exp2 = r; r += 1
    ws.cell(row=r, column=1, value="2차 광고"); r_ad2 = r; r += 1
    for a, b in labels:
        ws.cell(row=a, column=1, value=b)

    names = ["가", "나", "다"]
    yk = ["ㅇ", None, "o"]
    prods = ["수면", "", "스팟멜트크림"]
    status = ["미팅 대기", "1차 체험진행", "1차 광고완료"]
    acc = [D(2026, 9, 1), D(2026, 8, 1), D(2026, 7, 1)]
    ship = [D(2026, 9, 2), D(2026, 8, 2), D(2026, 7, 2)]   # 제품 발송일 — 제품으로 오인하면 안 됨
    ad1 = [None, None, D(2026, 8, 15)]
    exp2 = [None, D(2026, 9, 20), None]
    ad2 = [None, None, D(2026, 9, 30)]
    for i, n in enumerate(names):
        c = 3 + i
        ws.cell(row=r_name, column=c, value=n)
        if yk_row:
            ws.cell(row=9, column=c, value=yk[i])
        ws.cell(row=r_status, column=c, value=status[i])
        if product_row:
            ws.cell(row=r_prod, column=c, value=prods[i])
        ws.cell(row=r_acc, column=c, value=acc[i])
        ws.cell(row=r_ship, column=c, value=ship[i])
        ws.cell(row=r_ad1, column=c, value=ad1[i])
        ws.cell(row=r_exp2, column=c, value=exp2[i])
        ws.cell(row=r_ad2, column=c, value=ad2[i])
    wb.save(path)


def make_rawdata_wb(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw_Data"
    ws.append(["h"] * 21)
    def row(ono, dt, ytber, qty, product, amount, settle, cls="", opt="단품 1개", store="초방리농장"):
        r = [None] * 21
        r[0] = "x"; r[1] = ono; r[2] = dt; r[3] = "배송완료"; r[5] = ytber; r[6] = ""
        r[10] = opt; r[12] = qty; r[13] = "구매자"; r[15] = product; r[16] = store
        r[17] = amount; r[18] = settle; r[19] = "네이버"; r[20] = cls
        return r
    ws.append(row("o1", "2026-09-01 10:00", "가", 2, "흑염소", 0, 0))
    ws.append(row("o2", "2026-09-02 10:00", "나", 1, "흑염소", 0, 0))
    ws.append(row("o3", "2026-09-03 10:00", "가", 1, "화장품", 30000, 27000))
    ws.append(row("o4", "2026-09-04 10:00", "다", 1, "화장품", 30000, 27000, cls="체험단"))
    ws.append(row("o5", "2026-09-05 10:00", "나", 1, "수면영양제", 50000, 45000, cls="체험단", opt="1박스"))
    ws.append(row("o6", "2026-09-06 10:00", "가", 1, "수면영양제", 50000, 45000, cls="체험단", opt="1박스"))
    wb.save(path)


# ---------------------------------------------------------------- parse_mail
class TestParseMail(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_owner_mark_variants(self):
        for v in ["ㅇ", "o", "O", "○", " ㅇ ", "o\n"]:
            self.assertTrue(parse_mail.is_owner_mark(v), v)
        for v in [None, "", "x", "0", "ㅇㅇ", "oo"]:
            self.assertFalse(parse_mail.is_owner_mark(v), repr(v))

    def test_header_lookup_same_before_after_insertion(self):
        p_old = self.dir / "old.xlsx"; p_new = self.dir / "new.xlsx"
        make_mail_wb(p_old, MAIL_ROWS, with_owner=False)
        make_mail_wb(p_new, MAIL_ROWS, with_owner=True, owner_marks=[None] * len(MAIL_ROWS))
        a = parse_mail.parse_mail_workbook(p_old)
        b = parse_mail.parse_mail_workbook(p_new)
        self.assertEqual(a, b)
        for k, v in EXPECTED_ALL.items():
            self.assertEqual(a[k], v, k)
        self.assertEqual(a["by_month"]["2026-09"]["sent"], 3)
        self.assertEqual(a["by_month"]["2026-09"]["ad"], 1)      # 광고수락일 9월 귀속
        self.assertEqual(a["by_month"]["2026-08"]["exp"], 1)     # 검토 행 제외

    def test_owner_filter_counts_only_marked_rows(self):
        p = self.dir / "f.xlsx"
        make_mail_wb(p, MAIL_ROWS, with_owner=True, owner_marks=["ㅇ", None, "O", "x", "○", "ㅇ"])
        r = parse_mail.parse_mail_workbook(p, owner_filter=parse_mail.is_owner_mark)
        self.assertEqual(r["total_sent"], 3)
        self.assertEqual(r["replied"], 4)
        self.assertEqual(r["exp_total_approx"], 2)
        self.assertEqual(r["ad_total"], 1)
        self.assertEqual(r["etc_excluded"], 0)

    def test_owner_filter_without_owner_column_gives_zero(self):
        p = self.dir / "noowner.xlsx"
        make_mail_wb(p, MAIL_ROWS, with_owner=False)
        r = parse_mail.parse_mail_workbook(p, owner_filter=parse_mail.is_owner_mark)
        self.assertEqual(r["total_sent"], 0)
        self.assertEqual(r["by_month"], {})

    def test_sheet_name_suffix_allowed(self):
        p = self.dir / "beauty.xlsx"
        make_mail_wb(p, MAIL_ROWS, with_owner=False, sheet="메일발송현황(미니 대시보드 수식 수정 필요)")
        r = parse_mail.parse_mail_workbook(p)
        self.assertEqual(r["total_sent"], EXPECTED_ALL["total_sent"])

    def test_missing_required_header_raises(self):
        with self.assertRaises(ValueError):
            parse_mail.header_columns(["구분", "분류", "발송일"])

    def test_merge_results(self):
        p = self.dir / "m.xlsx"
        make_mail_wb(p, MAIL_ROWS, with_owner=False)
        r = parse_mail.parse_mail_workbook(p)
        m = parse_mail.merge_mail_results([r, r, parse_mail.empty_result()])
        self.assertEqual(m["total_sent"], 10)
        self.assertEqual(m["by_month"]["2026-09"]["sent"], 6)
        self.assertEqual(m["reply_rate"], round(10 / 10, 4))

    def test_find_shared_xlsx_latest_and_exclusions(self):
        for n in ["1. 뷰티 인플루언서 관리_공유_260701.xlsx",
                  "1. 뷰티 인플루언서 관리_공유_260615.xlsx",
                  "~$1. 뷰티 인플루언서 관리_공유_260701.xlsx",
                  "1. 뷰티 인플루언서 관리_공유_260701_백업_261007.xlsx",
                  "1. 뷰티 인플루언서 관리_공유_260801_backup.xlsx",
                  "1. 유튜브 인플루언서 관리_공유_260619.xlsx"]:
            (self.dir / n).write_bytes(b"")
        self.assertEqual(parse_mail.find_shared_xlsx("뷰티", self.dir).name,
                         "1. 뷰티 인플루언서 관리_공유_260701.xlsx")
        self.assertEqual(parse_mail.find_shared_xlsx("유튜브", self.dir).name,
                         "1. 유튜브 인플루언서 관리_공유_260619.xlsx")
        self.assertIsNone(parse_mail.find_shared_xlsx("슬립이지", self.dir))


# ---------------------------------------------------------------- parse_inf
class TestParseInf(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.cfg = {"name_map": {"다": "다다"}, "sponsor_extra": {}}

    def tearDown(self):
        self.tmp.cleanup()

    def test_label_lookup_survives_row_insertion(self):
        p1 = self.dir / "a.xlsx"; p2 = self.dir / "b.xlsx"
        make_master_wb(p1)
        make_master_wb(p2, extra_rows_before_name=3)
        a = parse_inf.parse_master(p1, self.cfg)
        b = parse_inf.parse_master(p2, self.cfg)
        for k in ("managed_set", "inf_status", "exp_by_month", "ad_by_month", "ad_total", "per_influencer",
                  "yk_marked", "product_by_name"):
            self.assertEqual(a[k], b[k], k)
        self.assertEqual(a["managed_set"], ["가", "나", "다"])
        self.assertEqual(a["ad_total"], 2)
        self.assertEqual(a["exp_by_month"], {"2026-07": 1, "2026-08": 1, "2026-09": 1})
        self.assertEqual(a["per_influencer"]["나"]["exp_count"], 2)   # 체험수락일 + 2차 체험
        self.assertEqual(a["inf_status"]["미팅_대기"], 1)

    def test_yk_marked_and_products(self):
        p = self.dir / "a.xlsx"
        make_master_wb(p)
        r = parse_inf.parse_master(p, self.cfg)
        self.assertEqual(r["yk_marked"], ["가", "다다"])              # name_map 적용
        self.assertEqual(r["product_by_name"], {"가": "수면영양제", "다다": "화장품"})  # 빈값은 제외

    def test_product_row_not_confused_with_shipping_date(self):
        p = self.dir / "a.xlsx"
        make_master_wb(p, product_row=False)
        r = parse_inf.parse_master(p, self.cfg)
        self.assertEqual(r["product_by_name"], {})   # '제품 발송일' 행을 제품으로 읽으면 안 됨

    def test_missing_yk_row(self):
        p = self.dir / "a.xlsx"
        make_master_wb(p, yk_row=False)
        r = parse_inf.parse_master(p, self.cfg)
        self.assertEqual(r["yk_marked"], [])
        self.assertEqual(r["managed_count"], 3)

    def test_normalize_product(self):
        self.assertEqual(parse_inf.normalize_product("수면"), "수면영양제")
        self.assertEqual(parse_inf.normalize_product("슬립이지"), "수면영양제")
        self.assertEqual(parse_inf.normalize_product("스팟멜트 크림"), "화장품")
        self.assertEqual(parse_inf.normalize_product("흑염소"), "흑염소")
        self.assertEqual(parse_inf.normalize_product(None), "")
        self.assertEqual(parse_inf.normalize_product("  "), "")

    def test_name_filter(self):
        p = self.dir / "a.xlsx"
        make_master_wb(p)
        r = parse_inf.parse_master(p, self.cfg, name_filter=lambda n: n in {"가", "다다"})
        self.assertEqual(r["managed_set"], ["가", "다"])
        self.assertEqual(set(r["per_influencer"]), {"가", "다다"})
        self.assertEqual(r["yk_marked"], ["가", "다다"])     # 필터와 무관하게 전체 수집
        self.assertEqual(r["ad_total"], 2)

    def test_build_row_index_first_occurrence(self):
        rows = [("항목",), ("현재 상태",), ("항목",), ("제품 발송일",), ("제품",)]
        idx = parse_inf.build_row_index(rows)
        self.assertEqual(idx["항목"], 0)
        self.assertEqual(idx["제품"], 4)
        self.assertEqual(idx["제품발송일"], 3)


# ---------------------------------------------------------------- build_snapshot name_filter
class TestSnapshotNameFilter(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.raw = self.dir / "raw.xlsx"
        make_rawdata_wb(self.raw)
        self._orig = build_snapshot.RAWDATA_PATH
        build_snapshot.RAWDATA_PATH = self.raw
        self.cfg = {
            "product_registry": {"default_product": "흑염소"},
            "general_sales_label": "기타/일반",
            "product_unit_price": 120000, "general_unit_price": 130000,
            "products": [{"key": "흑염소"}, {"key": "수면영양제"}, {"key": "화장품"}],
            "product_profit": {}, "product_units": {},
        }

    def tearDown(self):
        build_snapshot.RAWDATA_PATH = self._orig
        self.tmp.cleanup()

    def test_default_unfiltered(self):
        bp, _ = build_snapshot.aggregate_by_product_store("2026-09", self.cfg)
        self.assertEqual(bp["흑염소"]["qty"], 3)
        self.assertEqual(bp["화장품"]["order_count"], 2)

    def test_filtered(self):
        f = lambda n: n == "가"
        bp, _ = build_snapshot.aggregate_by_product_store("2026-09", self.cfg, name_filter=f)
        self.assertEqual(bp["흑염소"]["qty"], 2)
        self.assertEqual(bp["흑염소"]["gross_revenue"], 240000)
        self.assertEqual(bp["화장품"]["order_count"], 1)
        self.assertEqual(bp["수면영양제"]["trial_units"], 1)
        cb = build_snapshot.aggregate_cosmetics_breakdown(None, name_filter=f)
        self.assertEqual(cb["trial"]["orders"], 0)
        cb_all = build_snapshot.aggregate_cosmetics_breakdown(None)
        self.assertEqual(cb_all["trial"]["orders"], 1)
        tr = build_snapshot.aggregate_product_trials(None, name_filter=f)
        self.assertEqual(tr["수면영양제"]["orders"], 1)
        tr_all = build_snapshot.aggregate_product_trials(None)
        self.assertEqual(tr_all["수면영양제"]["orders"], 2)


# ---------------------------------------------------------------- build_yk helpers
class TestBuildYkHelpers(unittest.TestCase):
    def setUp(self):
        import build_yk
        self.m = build_yk

    def test_mail_cum_through(self):
        bm = {"2026-08": {"sent": 2, "replied": 1, "meeting": 1, "exp": 0, "ad": 0},
              "2026-09": {"sent": 3, "replied": 1, "meeting": 0, "exp": 2, "ad": 1},
              "2026-10": {"sent": 1, "replied": 0, "meeting": 0, "exp": 0, "ad": 0}}
        c = self.m.mail_cum_through(bm, "2026-09")
        self.assertEqual((c["total_sent"], c["replied"], c["meeting_total"], c["exp_total"], c["ad_total"]),
                         (5, 2, 1, 2, 1))
        self.assertEqual(c["reply_rate"], round(2 / 5, 4))
        self.assertEqual(self.m.mail_cum_through({}, "2026-09")["total_sent"], 0)

    def test_filter_settlement(self):
        s = {"settlement_month": "2026-09", "summaries": [
            {"ytber": "가", "qty": 1}, {"ytber": "나", "qty": 2}, {"ytber": "기타/일반", "qty": 9, "is_general": True}],
            "unregistered": [{"name": "나"}, {"name": "엑스"}]}
        f = self.m.filter_settlement(s, {"가"})
        self.assertEqual([x["ytber"] for x in f["summaries"]], ["가"])
        self.assertEqual(f["unregistered"], [])
        self.assertEqual(f["settlement_month"], "2026-09")

    def test_resolve_products(self):
        names = ["가", "나", "다"]
        got = self.m.resolve_products(names, {"가": "수면영양제"}, {"나": "화장품"})
        self.assertEqual(got, {"가": "수면영양제", "나": "화장품", "다": "미지정"})


if __name__ == "__main__":
    unittest.main()
